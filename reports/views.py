from django.http import HttpResponse
from django.contrib.auth.decorators import login_required

from openpyxl import Workbook

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER

from students.models import Student
from attendance.models import Attendance
from marks.models import Marks
from faculty.models import Faculty


def get_faculty(request):
    """
    Looks up the Faculty profile linked to the logged-in user.
    Prefers the direct user link; falls back to matching by
    email for older records that predate that link.
    """

    faculty = Faculty.objects.filter(
        user=request.user
    ).first()

    if faculty:
        return faculty

    return Faculty.objects.filter(
        email=request.user.email
    ).first()


def grade_for_percentage(percentage):

    if percentage >= 90:
        return "A+"

    elif percentage >= 80:
        return "A"

    elif percentage >= 70:
        return "B"

    elif percentage >= 60:
        return "C"

    elif percentage >= 50:
        return "D"

    else:
        return "F"


def pdf_document(response, subtitle):
    """
    Sets up a SimpleDocTemplate with the standard SmartCollege
    title block, returns (document, elements, styles).
    """

    document = SimpleDocTemplate(
        response,
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]
    title_style.alignment = TA_CENTER

    elements = []

    elements.append(
        Paragraph(
            "SMART COLLEGE",
            title_style
        )
    )

    elements.append(
        Paragraph(
            subtitle,
            styles["Heading2"]
        )
    )

    elements.append(
        Spacer(1, 15)
    )

    return document, elements, styles


def styled_table(table_data, col_widths):

    table = Table(
        table_data,
        colWidths=col_widths
    )

    table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1f2937")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "ALIGN",
                (1, 0),
                (-1, -1),
                "CENTER"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),

            (
                "FONTNAME",
                (0, 1),
                (-1, -1),
                "Helvetica"
            ),

            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                9
            ),

            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                6
            ),

            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
        ])
    )

    return table


# =========================================================
# STUDENT ATTENDANCE EXCEL
# =========================================================

@login_required(login_url="/accounts/login/")
def student_attendance_excel(request):

    student = Student.objects.filter(
        user=request.user
    ).first()

    if student is None:
        return HttpResponse(
            "Student profile not found.",
            status=404
        )

    attendance_records = Attendance.objects.filter(
        student=student
    ).select_related(
        "subject"
    ).order_by(
        "-date"
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Attendance"

    worksheet.append([
        "Student Name",
        "Roll No",
        "Subject",
        "Subject Code",
        "Date",
        "Status"
    ])

    for record in attendance_records:

        status = (
            "Present"
            if record.present
            else "Absent"
        )

        worksheet.append([
            student.name,
            student.roll_no,
            record.subject.name,
            record.subject.code,
            record.date,
            status
        ])

    worksheet.column_dimensions["A"].width = 25
    worksheet.column_dimensions["B"].width = 15
    worksheet.column_dimensions["C"].width = 25
    worksheet.column_dimensions["D"].width = 15
    worksheet.column_dimensions["E"].width = 15
    worksheet.column_dimensions["F"].width = 15

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="student_attendance.xlsx"'
    )

    workbook.save(response)

    return response


# =========================================================
# STUDENT MARKS EXCEL
# =========================================================

@login_required(login_url="/accounts/login/")
def student_marks_excel(request):

    student = Student.objects.filter(
        user=request.user
    ).first()

    if student is None:
        return HttpResponse(
            "Student profile not found.",
            status=404
        )

    marks_records = Marks.objects.filter(
        student=student
    ).select_related(
        "subject"
    ).order_by(
        "subject__name"
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Marks"

    worksheet.append([
        "Student Name",
        "Roll No",
        "Subject",
        "Subject Code",
        "CAT 1 Exam",
        "CAT 1 Assignment",
        "CAT 2 Exam",
        "CAT 2 Assignment",
        "CAT 3 Exam",
        "CAT 3 Assignment",
        "Total",
        "Average",
        "Percentage",
        "Grade"
    ])

    for mark in marks_records:

        total = mark.grand_total

        average = round(
            total / 3,
            2
        )

        percentage = round(
            (total / 150) * 100,
            2
        )

        if percentage >= 90:
            grade = "A+"

        elif percentage >= 80:
            grade = "A"

        elif percentage >= 70:
            grade = "B"

        elif percentage >= 60:
            grade = "C"

        elif percentage >= 50:
            grade = "D"

        else:
            grade = "F"

        worksheet.append([
            student.name,
            student.roll_no,
            mark.subject.name,
            mark.subject.code,
            mark.cat_1,
            mark.cat_1_assignment,
            mark.cat_2,
            mark.cat_2_assignment,
            mark.cat_3,
            mark.cat_3_assignment,
            total,
            average,
            percentage,
            grade
        ])

    worksheet.column_dimensions["A"].width = 25
    worksheet.column_dimensions["B"].width = 15
    worksheet.column_dimensions["C"].width = 25
    worksheet.column_dimensions["D"].width = 15
    worksheet.column_dimensions["E"].width = 10
    worksheet.column_dimensions["F"].width = 10
    worksheet.column_dimensions["G"].width = 10
    worksheet.column_dimensions["H"].width = 10
    worksheet.column_dimensions["I"].width = 12
    worksheet.column_dimensions["J"].width = 15
    worksheet.column_dimensions["K"].width = 10

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="student_marks.xlsx"'
    )

    workbook.save(response)

    return response


# =========================================================
# FACULTY ATTENDANCE EXCEL
# =========================================================

@login_required(login_url="/accounts/login/")
def faculty_attendance_excel(request):

    faculty = get_faculty(request)

    if faculty is None:
        return HttpResponse(
            "Faculty profile not found.",
            status=404
        )

    attendance_records = Attendance.objects.filter(
        subject__faculty=faculty
    ).select_related(
        "student",
        "subject"
    ).order_by(
        "-date",
        "student__name"
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Attendance"

    worksheet.append([
        "Student Name",
        "Roll No",
        "Subject",
        "Subject Code",
        "Date",
        "Status"
    ])

    for record in attendance_records:

        status = (
            "Present"
            if record.present
            else "Absent"
        )

        worksheet.append([
            record.student.name,
            record.student.roll_no,
            record.subject.name,
            record.subject.code,
            record.date,
            status
        ])

    worksheet.column_dimensions["A"].width = 25
    worksheet.column_dimensions["B"].width = 15
    worksheet.column_dimensions["C"].width = 25
    worksheet.column_dimensions["D"].width = 15
    worksheet.column_dimensions["E"].width = 15
    worksheet.column_dimensions["F"].width = 15

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="faculty_attendance.xlsx"'
    )

    workbook.save(response)

    return response


# =========================================================
# FACULTY MARKS EXCEL
# =========================================================

@login_required(login_url="/accounts/login/")
def faculty_marks_excel(request):

    faculty = get_faculty(request)

    if faculty is None:
        return HttpResponse(
            "Faculty profile not found.",
            status=404
        )

    marks_records = Marks.objects.filter(
        subject__faculty=faculty
    ).select_related(
        "student",
        "subject"
    ).order_by(
        "student__name",
        "subject__name"
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Marks"

    worksheet.append([
        "Student Name",
        "Roll No",
        "Subject",
        "Subject Code",
        "CAT 1 Exam",
        "CAT 1 Assignment",
        "CAT 2 Exam",
        "CAT 2 Assignment",
        "CAT 3 Exam",
        "CAT 3 Assignment",
        "Total",
        "Average",
        "Percentage",
        "Grade"
    ])

    for mark in marks_records:

        total = mark.grand_total

        average = round(
            total / 3,
            2
        )

        percentage = round(
            (total / 150) * 100,
            2
        )

        if percentage >= 90:
            grade = "A+"

        elif percentage >= 80:
            grade = "A"

        elif percentage >= 70:
            grade = "B"

        elif percentage >= 60:
            grade = "C"

        elif percentage >= 50:
            grade = "D"

        else:
            grade = "F"

        worksheet.append([
            mark.student.name,
            mark.student.roll_no,
            mark.subject.name,
            mark.subject.code,
            mark.cat_1,
            mark.cat_1_assignment,
            mark.cat_2,
            mark.cat_2_assignment,
            mark.cat_3,
            mark.cat_3_assignment,
            total,
            average,
            percentage,
            grade
        ])

    worksheet.column_dimensions["A"].width = 25
    worksheet.column_dimensions["B"].width = 15
    worksheet.column_dimensions["C"].width = 25
    worksheet.column_dimensions["D"].width = 15
    worksheet.column_dimensions["E"].width = 10
    worksheet.column_dimensions["F"].width = 10
    worksheet.column_dimensions["G"].width = 10
    worksheet.column_dimensions["H"].width = 10
    worksheet.column_dimensions["I"].width = 12
    worksheet.column_dimensions["J"].width = 15
    worksheet.column_dimensions["K"].width = 10

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="faculty_marks.xlsx"'
    )

    workbook.save(response)

    return response


# =========================================================
# STUDENT ATTENDANCE PDF
# =========================================================

@login_required(login_url="/accounts/login/")
def student_attendance_pdf(request):

    student = Student.objects.filter(
        user=request.user
    ).first()

    if student is None:
        return HttpResponse(
            "Student profile not found.",
            status=404
        )

    attendance_records = Attendance.objects.filter(
        student=student
    ).select_related(
        "subject"
    ).order_by(
        "-date"
    )

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="student_attendance.pdf"'
    )

    document = SimpleDocTemplate(
        response,
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]
    title_style.alignment = TA_CENTER

    elements = []

    # ---------------------------------------------------------
    # TITLE
    # ---------------------------------------------------------

    elements.append(
        Paragraph(
            "SMART COLLEGE",
            title_style
        )
    )

    elements.append(
        Paragraph(
            "Student Attendance Report",
            styles["Heading2"]
        )
    )

    elements.append(
        Spacer(1, 15)
    )

    # ---------------------------------------------------------
    # STUDENT DETAILS
    # ---------------------------------------------------------

    elements.append(
        Paragraph(
            f"<b>Student Name:</b> {student.name}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            f"<b>Roll No:</b> {student.roll_no}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            f"<b>Department:</b> "
            f"{student.department.name}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(1, 15)
    )

    # ---------------------------------------------------------
    # TABLE DATA
    # ---------------------------------------------------------

    table_data = [
        [
            "Subject",
            "Code",
            "Date",
            "Status"
        ]
    ]

    total = 0
    present = 0

    for record in attendance_records:

        total += 1

        if record.present:

            status = "Present"
            present += 1

        else:

            status = "Absent"

        table_data.append([
            record.subject.name,
            record.subject.code,
            record.date.strftime("%d-%m-%Y"),
            status
        ])

    # ---------------------------------------------------------
    # ATTENDANCE PERCENTAGE
    # ---------------------------------------------------------

    if total > 0:

        attendance_percentage = round(
            (present / total) * 100,
            2
        )

    else:

        attendance_percentage = 0

    elements.append(
        Paragraph(
            f"<b>Total Classes:</b> {total} "
            f"&nbsp;&nbsp;&nbsp; "
            f"<b>Present:</b> {present} "
            f"&nbsp;&nbsp;&nbsp; "
            f"<b>Attendance:</b> "
            f"{attendance_percentage}%",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(1, 15)
    )

    # ---------------------------------------------------------
    # ATTENDANCE TABLE
    # ---------------------------------------------------------

    table = Table(
        table_data,
        colWidths=[
            190,
            70,
            90,
            70
        ]
    )

    table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1f2937")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "ALIGN",
                (1, 0),
                (-1, -1),
                "CENTER"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),

            (
                "FONTNAME",
                (0, 1),
                (-1, -1),
                "Helvetica"
            ),

            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                9
            ),

            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                6
            ),

            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                6
            ),
        ])
    )

    elements.append(table)

    # ---------------------------------------------------------
    # BUILD PDF
    # ---------------------------------------------------------

    document.build(elements)

    return response


# =========================================================
# STUDENT MARKS PDF
# =========================================================

@login_required(login_url="/accounts/login/")
def student_marks_pdf(request):

    student = Student.objects.filter(
        user=request.user
    ).first()

    if student is None:
        return HttpResponse(
            "Student profile not found.",
            status=404
        )

    marks_records = Marks.objects.filter(
        student=student
    ).select_related(
        "subject"
    ).order_by(
        "subject__name"
    )

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="student_marks.pdf"'
    )

    document, elements, styles = pdf_document(
        response,
        "Student Marks Report"
    )

    elements.append(
        Paragraph(
            f"<b>Student Name:</b> {student.name}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            f"<b>Roll No:</b> {student.roll_no}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            f"<b>Department:</b> "
            f"{student.department.name}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(1, 15)
    )

    table_data = [
        [
            "Subject",
            "Code",
            "CAT1 E",
            "CAT1 A",
            "CAT2 E",
            "CAT2 A",
            "CAT3 E",
            "CAT3 A",
            "Total",
            "%",
            "Grade"
        ]
    ]

    for mark in marks_records:

        total = mark.grand_total

        percentage = round(
            (total / 150) * 100,
            2
        )

        grade = grade_for_percentage(percentage)

        table_data.append([
            mark.subject.name,
            mark.subject.code,
            mark.cat_1,
            mark.cat_1_assignment,
            mark.cat_2,
            mark.cat_2_assignment,
            mark.cat_3,
            mark.cat_3_assignment,
            total,
            f"{percentage}%",
            grade
        ])

    table = styled_table(
        table_data,
        col_widths=[105, 45, 35, 35, 35, 35, 35, 35, 45, 45, 40]
    )

    elements.append(table)

    document.build(elements)

    return response


# =========================================================
# FACULTY ATTENDANCE PDF
# =========================================================

@login_required(login_url="/accounts/login/")
def faculty_attendance_pdf(request):

    faculty = get_faculty(request)

    if faculty is None:
        return HttpResponse(
            "Faculty profile not found.",
            status=404
        )

    attendance_records = Attendance.objects.filter(
        subject__faculty=faculty
    ).select_related(
        "student",
        "subject"
    ).order_by(
        "-date",
        "student__name"
    )

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="faculty_attendance.pdf"'
    )

    document, elements, styles = pdf_document(
        response,
        "Faculty Attendance Report"
    )

    elements.append(
        Paragraph(
            f"<b>Faculty Name:</b> {faculty.name}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            f"<b>Department:</b> "
            f"{faculty.department.name}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(1, 15)
    )

    table_data = [
        [
            "Student",
            "Roll No",
            "Subject",
            "Date",
            "Status"
        ]
    ]

    for record in attendance_records:

        status = (
            "Present"
            if record.present
            else "Absent"
        )

        table_data.append([
            record.student.name,
            record.student.roll_no,
            record.subject.name,
            record.date.strftime("%d-%m-%Y"),
            status
        ])

    table = styled_table(
        table_data,
        col_widths=[130, 70, 130, 80, 65]
    )

    elements.append(table)

    document.build(elements)

    return response


# =========================================================
# FACULTY MARKS PDF
# =========================================================

@login_required(login_url="/accounts/login/")
def faculty_marks_pdf(request):

    faculty = get_faculty(request)

    if faculty is None:
        return HttpResponse(
            "Faculty profile not found.",
            status=404
        )

    marks_records = Marks.objects.filter(
        subject__faculty=faculty
    ).select_related(
        "student",
        "subject"
    ).order_by(
        "student__name",
        "subject__name"
    )

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="faculty_marks.pdf"'
    )

    document, elements, styles = pdf_document(
        response,
        "Faculty Marks Report"
    )

    elements.append(
        Paragraph(
            f"<b>Faculty Name:</b> {faculty.name}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            f"<b>Department:</b> "
            f"{faculty.department.name}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(1, 15)
    )

    table_data = [
        [
            "Student",
            "Roll No",
            "Subject",
            "Total",
            "%",
            "Grade"
        ]
    ]

    for mark in marks_records:

        total = mark.grand_total

        percentage = round(
            (total / 150) * 100,
            2
        )

        grade = grade_for_percentage(percentage)

        table_data.append([
            mark.student.name,
            mark.student.roll_no,
            mark.subject.name,
            total,
            f"{percentage}%",
            grade
        ])

    table = styled_table(
        table_data,
        col_widths=[120, 70, 130, 55, 55, 45]
    )

    elements.append(table)

    document.build(elements)

    return response


# =========================================================
# ADMIN — STUDENTS EXCEL (respects School/Degree/Class filters)
# =========================================================

@login_required(login_url="/accounts/login/")
def admin_students_excel(request):

    if not request.user.is_superuser:
        return HttpResponse("Admins only.", status=403)

    from django.db.models import Q
    from accounts import scope as sc

    scope = sc.parse_scope(request.GET)

    students = sc.scoped_students(scope).order_by(
        "department__name", "course__name", "year", "name"
    )

    q = request.GET.get("q", "").strip()

    if q:
        students = students.filter(
            Q(name__icontains=q)
            | Q(roll_no__icontains=q)
            | Q(sif_number__icontains=q)
            | Q(email__icontains=q)
        )

    students = list(students)

    stats = sc.per_student_stats([s.id for s in students])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"

    sheet.append([
        "Name", "Roll No", "SIF Number", "Email",
        "School", "Degree", "Year",
        "Attendance %", "Average Marks %",
    ])

    for s in students:

        sheet.append([
            s.name,
            s.roll_no,
            s.sif_number or "",
            s.email,
            s.department.name,
            s.course.name if s.course else "",
            s.year,
            stats[s.id]["attendance_pct"],
            stats[s.id]["marks_pct"],
        ])

    for column, width in zip("ABCDEFGHI", [24, 12, 16, 30, 32, 42, 7, 14, 16]):
        sheet.column_dimensions[column].width = width

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="students.xlsx"'
    )

    workbook.save(response)

    return response
