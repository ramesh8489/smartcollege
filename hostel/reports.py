import io
from decimal import Decimal
from django.http import HttpResponse
from django.shortcuts import render
from django.views import View
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied
from django.utils import timezone
from django.db.models import Count, Q, Sum

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

from .models import (
    Hostel,
    HostelRoom,
    HostelBed,
    HostelAllocation,
    HostelAttendance,
    HostelComplaint,
    HostelVisitor,
    HostelApplication,
    HostelCheckInOut,
    HostelRoomTransfer,
)
from fees.models import FeeRecord
from .views import get_user_assigned_hostel_ids, _is_hostel_admin


def _apply_excel_header_style(ws, headers):
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 26


def _autofit_excel_columns(ws, min_width=12, max_width=45):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 3, max_width))


def _build_pdf_report(title, headers, data, col_widths=None, orientation="portrait", subtitle=""):
    buffer = io.BytesIO()
    pagesize = landscape(A4) if orientation == "landscape" else A4
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=15,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=3,
        alignment=TA_CENTER,
    )
    sub_style = ParagraphStyle(
        "DocSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        textColor=colors.HexColor("#475569"),
        spaceAfter=12,
        alignment=TA_CENTER,
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9.5,
    )
    cell_header = ParagraphStyle(
        "TableHead",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=TA_CENTER,
    )

    elements = [
        Paragraph("SMARTCOLLEGE INSTITUTIONAL HOSTEL REPORT", title_style),
        Paragraph(
            f"<b>{title}</b> &middot; {subtitle or 'Official Audit Roster'} &middot; Generated: {timezone.now().strftime('%d %b %Y %H:%M')}",
            sub_style,
        ),
    ]

    table_data = [[Paragraph(h, cell_header) for h in headers]]
    for row in data:
        row_cells = []
        for c in row:
            val = str(c) if c is not None else ""
            row_cells.append(Paragraph(val, cell_style))
        table_data.append(row_cells)

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    t_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    for i in range(1, len(table_data)):
        if i % 2 == 0:
            t_style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#f8fafc")))
    table.setStyle(TableStyle(t_style))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


class HostelReportsHubView(LoginRequiredMixin, View):
    """
    Central reporting hub offering filtered report exports for Occupancy, Allocations,
    Room Availability, Attendance, Fees, Complaints, and Visitor Logs in both Excel & PDF formats.
    """
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_hostels = Hostel.objects.filter(
            id__in=get_user_assigned_hostel_ids(request.user),
            is_active=True,
        )

        reports_list = [
            {
                "id": "occupancy",
                "name": "Hostel Occupancy & Capacity",
                "desc": "Room-wise and bed-level occupancy breakdown, capacity utilization, and vacancies.",
                "excel_url": "hostel:export_occupancy_excel",
                "pdf_url": "hostel:export_occupancy_pdf",
                "badge": "Capacity & Intake",
            },
            {
                "id": "allocations",
                "name": "Resident Allocation Roster",
                "desc": "Active student room allotments, roll numbers, admission details, and expected checkouts.",
                "excel_url": "hostel:export_allocations_excel",
                "pdf_url": "hostel:export_allocations_pdf",
                "badge": "Residents",
            },
            {
                "id": "availability",
                "name": "Vacant Bed & Room Availability",
                "desc": "Allocatable rooms with vacant beds ready for upcoming academic intake.",
                "excel_url": "hostel:export_availability_excel",
                "pdf_url": "hostel:export_availability_pdf",
                "badge": "Allotment",
            },
            {
                "id": "attendance",
                "name": "Hostel Roll Call Attendance",
                "desc": "Nightly curfew attendance logs, resident present/absent roll calls, and leave tracking.",
                "excel_url": "hostel:export_attendance_excel",
                "pdf_url": "hostel:export_attendance_pdf",
                "badge": "Discipline",
            },
            {
                "id": "fees",
                "name": "Hostel Fees & Outstanding Dues",
                "desc": "Hostel rent, mess fees, security deposits, collection status, and overdue amounts.",
                "excel_url": "hostel:export_fees_excel",
                "pdf_url": "hostel:export_fees_pdf",
                "badge": "Finance",
            },
            {
                "id": "complaints",
                "name": "Maintenance & Complaints Log",
                "desc": "Electrical, plumbing, carpentry, and WiFi tickets with staff assignment and resolution.",
                "excel_url": "hostel:export_complaints_excel",
                "pdf_url": "hostel:export_complaints_pdf",
                "badge": "Facilities",
            },
            {
                "id": "applications",
                "name": "Hostel Admission & Application Log",
                "desc": "Student applications, room preferences, review decisions, and waitlist records.",
                "excel_url": "hostel:export_applications_excel",
                "pdf_url": "hostel:export_applications_pdf",
                "badge": "Admissions",
            },
            {
                "id": "checkinout",
                "name": "Check-in & Check-out History",
                "desc": "Resident lifecycle history, move-in verifications, and vacate/clearance records.",
                "excel_url": "hostel:export_checkinout_excel",
                "pdf_url": "hostel:export_checkinout_pdf",
                "badge": "Lifecycle",
            },
            {
                "id": "transfers",
                "name": "Room & Bed Transfer Audit Trail",
                "desc": "Resident room reassignments, inter-hostel transfers, warden approvals, and bed releases.",
                "excel_url": "hostel:export_transfers_excel",
                "pdf_url": "hostel:export_transfers_pdf",
                "badge": "Transfers",
            },
            {
                "id": "visitors",
                "name": "Visitor Passes & Security Gate Log",
                "desc": "Campus visitor clearances, relationship records, and check-in / check-out timestamps.",
                "excel_url": "hostel:export_visitors_excel",
                "pdf_url": "hostel:export_visitors_pdf",
                "badge": "Security",
            },
        ]

        return render(
            request,
            "hostel/admin_reports_hub.html",
            {
                "hostels": assigned_hostels,
                "reports": reports_list,
            },
        )


# ==============================================================================
# 1. OCCUPANCY REPORT
# ==============================================================================
class ExportHostelOccupancyExcelView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        rooms_qs = HostelRoom.objects.filter(hostel_id__in=assigned_ids, is_active=True).select_related(
            "hostel", "block", "floor"
        ).prefetch_related("beds")

        if hostel_id and hostel_id.isdigit():
            rooms_qs = rooms_qs.filter(hostel_id=int(hostel_id))

        rooms_qs = rooms_qs.order_by("hostel__code", "block__name", "floor__floor_number", "room_number")

        wb = Workbook()
        ws = wb.active
        ws.title = "Hostel Occupancy"

        headers = [
            "Hostel Code",
            "Hostel Name",
            "Block",
            "Floor",
            "Room No",
            "Room Type",
            "Capacity",
            "Occupied Beds",
            "Available Beds",
            "Maintenance Beds",
            "Occupancy %",
            "Room Status",
        ]
        _apply_excel_header_style(ws, headers)

        for r in rooms_qs:
            beds = list(r.beds.filter(is_active=True))
            total = len(beds) or r.capacity
            occ = sum(1 for b in beds if b.status == "OCCUPIED")
            avail = sum(1 for b in beds if b.status == "AVAILABLE")
            maint = sum(1 for b in beds if b.status == "MAINTENANCE")
            occ_pct = f"{round((occ / total) * 100, 1)}%" if total > 0 else "0%"

            ws.append([
                r.hostel.code,
                r.hostel.name,
                r.block.name,
                r.floor.name,
                r.room_number,
                r.get_room_type_display(),
                r.capacity,
                occ,
                avail,
                maint,
                occ_pct,
                r.get_status_display(),
            ])

        _autofit_excel_columns(ws)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="Hostel_Occupancy_Report.xlsx"'
        wb.save(response)
        return response


class ExportHostelOccupancyPdfView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        rooms_qs = HostelRoom.objects.filter(hostel_id__in=assigned_ids, is_active=True).select_related(
            "hostel", "block", "floor"
        ).prefetch_related("beds")

        if hostel_id and hostel_id.isdigit():
            rooms_qs = rooms_qs.filter(hostel_id=int(hostel_id))

        rooms_qs = rooms_qs.order_by("hostel__code", "block__name", "floor__floor_number", "room_number")

        headers = ["Hostel", "Block", "Floor", "Room", "Type", "Cap", "Occ", "Avail", "Maint", "Rate", "Status"]
        data = []
        for r in rooms_qs:
            beds = list(r.beds.filter(is_active=True))
            total = len(beds) or r.capacity
            occ = sum(1 for b in beds if b.status == "OCCUPIED")
            avail = sum(1 for b in beds if b.status == "AVAILABLE")
            maint = sum(1 for b in beds if b.status == "MAINTENANCE")
            rate = f"{round((occ / total) * 100)}%" if total > 0 else "0%"
            data.append([
                r.hostel.code,
                r.block.name,
                r.floor.name,
                r.room_number,
                r.get_room_type_display()[:6],
                r.capacity,
                occ,
                avail,
                maint,
                rate,
                r.get_status_display(),
            ])

        pdf_bytes = _build_pdf_report(
            "Hostel Occupancy & Capacity Report",
            headers,
            data,
            orientation="landscape",
        )
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="Hostel_Occupancy_Report.pdf"'
        return response


# ==============================================================================
# 2. RESIDENT ALLOCATIONS ROSTER
# ==============================================================================
class ExportHostelAllocationsExcelView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        allocs = HostelAllocation.objects.filter(hostel_id__in=assigned_ids).select_related(
            "student", "student__department", "hostel", "block", "floor", "room", "bed"
        )
        if hostel_id and hostel_id.isdigit():
            allocs = allocs.filter(hostel_id=int(hostel_id))

        allocs = allocs.order_by("-allocation_date", "-id")

        wb = Workbook()
        ws = wb.active
        ws.title = "Allocations Roster"

        headers = [
            "Allocation ID",
            "Student Name",
            "Roll No",
            "Department",
            "Hostel Code",
            "Hostel Name",
            "Block",
            "Room No",
            "Bed No",
            "Academic Year",
            "Allocated Date",
            "Expected Checkout",
            "Status",
        ]
        _apply_excel_header_style(ws, headers)

        for a in allocs:
            dept_name = a.student.department.name if a.student.department else "—"
            ws.append([
                a.allocation_id,
                a.student.name,
                a.student.roll_no,
                dept_name,
                a.hostel.code,
                a.hostel.name,
                a.block.name if a.block else "—",
                a.room.room_number if a.room else "—",
                a.bed.bed_number if a.bed else "—",
                a.academic_year,
                a.allocation_date.strftime("%d-%m-%Y") if a.allocation_date else "—",
                a.expected_checkout_date.strftime("%d-%m-%Y") if a.expected_checkout_date else "—",
                a.get_status_display(),
            ])

        _autofit_excel_columns(ws)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="Hostel_Allocations_Roster.xlsx"'
        wb.save(response)
        return response


class ExportHostelAllocationsPdfView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        allocs = HostelAllocation.objects.filter(hostel_id__in=assigned_ids).select_related(
            "student", "student__department", "hostel", "block", "room", "bed"
        )
        if hostel_id and hostel_id.isdigit():
            allocs = allocs.filter(hostel_id=int(hostel_id))

        allocs = allocs.order_by("-allocation_date", "-id")

        headers = ["Alloc ID", "Student Name", "Roll No", "Hostel", "Room", "Bed", "Alloc Date", "Exp Checkout", "Status"]
        data = []
        for a in allocs:
            data.append([
                a.allocation_id,
                a.student.name,
                a.student.roll_no,
                a.hostel.code,
                a.room.room_number if a.room else "—",
                a.bed.bed_number if a.bed else "—",
                a.allocation_date.strftime("%d-%b-%y") if a.allocation_date else "—",
                a.expected_checkout_date.strftime("%d-%b-%y") if a.expected_checkout_date else "—",
                a.get_status_display(),
            ])

        pdf_bytes = _build_pdf_report(
            "Resident Allotment Registry",
            headers,
            data,
            orientation="landscape",
        )
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="Hostel_Allocations_Roster.pdf"'
        return response


# ==============================================================================
# 3. VACANT BED & ROOM AVAILABILITY REPORT
# ==============================================================================
class ExportHostelAvailabilityExcelView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        rooms = (
            HostelRoom.objects.filter(hostel_id__in=assigned_ids, is_active=True, status="AVAILABLE")
            .select_related("hostel", "block", "floor")
            .prefetch_related("beds")
            .order_by("hostel__code", "block__name", "room_number")
        )
        if hostel_id and hostel_id.isdigit():
            rooms = rooms.filter(hostel_id=int(hostel_id))

        wb = Workbook()
        ws = wb.active
        ws.title = "Room Availability"

        headers = [
            "Hostel Code",
            "Hostel Name",
            "Block",
            "Floor",
            "Room No",
            "Room Type",
            "Capacity",
            "Vacant Beds Count",
            "Available Bed Numbers",
        ]
        _apply_excel_header_style(ws, headers)

        for r in rooms:
            avail_beds = [b.bed_number for b in r.beds.filter(status="AVAILABLE", is_active=True)]
            if avail_beds:
                ws.append([
                    r.hostel.code,
                    r.hostel.name,
                    r.block.name,
                    r.floor.name,
                    r.room_number,
                    r.get_room_type_display(),
                    r.capacity,
                    len(avail_beds),
                    ", ".join(avail_beds),
                ])

        _autofit_excel_columns(ws)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="Hostel_Room_Availability.xlsx"'
        wb.save(response)
        return response


class ExportHostelAvailabilityPdfView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        rooms = (
            HostelRoom.objects.filter(hostel_id__in=assigned_ids, is_active=True, status="AVAILABLE")
            .select_related("hostel", "block", "floor")
            .prefetch_related("beds")
            .order_by("hostel__code", "block__name", "room_number")
        )
        if hostel_id and hostel_id.isdigit():
            rooms = rooms.filter(hostel_id=int(hostel_id))

        headers = ["Hostel", "Block", "Floor", "Room", "Type", "Capacity", "Vacant Beds", "Allocatable Bed List"]
        data = []
        for r in rooms:
            avail_beds = [b.bed_number for b in r.beds.filter(status="AVAILABLE", is_active=True)]
            if avail_beds:
                data.append([
                    r.hostel.code,
                    r.block.name,
                    r.floor.name,
                    r.room_number,
                    r.get_room_type_display(),
                    r.capacity,
                    len(avail_beds),
                    ", ".join(avail_beds),
                ])

        pdf_bytes = _build_pdf_report(
            "Hostel Room & Bed Vacancy Roster",
            headers,
            data,
            orientation="portrait",
        )
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="Hostel_Room_Availability.pdf"'
        return response


# ==============================================================================
# 4. ATTENDANCE SUMMARY REPORT
# ==============================================================================
class ExportHostelAttendanceExcelView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        date_str = request.GET.get("date")

        att_qs = HostelAttendance.objects.filter(hostel_id__in=assigned_ids).select_related(
            "student", "hostel", "room", "marked_by"
        )
        if hostel_id and hostel_id.isdigit():
            att_qs = att_qs.filter(hostel_id=int(hostel_id))
        if date_str:
            att_qs = att_qs.filter(date=date_str)

        att_qs = att_qs.order_by("-date", "hostel__code", "room__room_number")

        wb = Workbook()
        ws = wb.active
        ws.title = "Hostel Attendance"

        headers = [
            "Roll Call Date",
            "Hostel Code",
            "Student Name",
            "Roll No",
            "Room No",
            "Status",
            "Remarks",
            "Marked By",
        ]
        _apply_excel_header_style(ws, headers)

        for a in att_qs:
            marker = a.marked_by.get_full_name() or a.marked_by.username if a.marked_by else "—"
            ws.append([
                a.date.strftime("%d-%m-%Y"),
                a.hostel.code,
                a.student.name,
                a.student.roll_no,
                a.room.room_number if a.room else "—",
                a.get_status_display(),
                a.remarks or "",
                marker,
            ])

        _autofit_excel_columns(ws)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="Hostel_Attendance_Report.xlsx"'
        wb.save(response)
        return response


class ExportHostelAttendancePdfView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        date_str = request.GET.get("date")

        att_qs = HostelAttendance.objects.filter(hostel_id__in=assigned_ids).select_related(
            "student", "hostel", "room", "marked_by"
        )
        if hostel_id and hostel_id.isdigit():
            att_qs = att_qs.filter(hostel_id=int(hostel_id))
        if date_str:
            att_qs = att_qs.filter(date=date_str)

        att_qs = att_qs.order_by("-date", "hostel__code", "room__room_number")

        headers = ["Date", "Hostel", "Student Name", "Roll No", "Room", "Status", "Remarks"]
        data = []
        for a in att_qs:
            data.append([
                a.date.strftime("%d-%b-%y"),
                a.hostel.code,
                a.student.name,
                a.student.roll_no,
                a.room.room_number if a.room else "—",
                a.get_status_display(),
                (a.remarks or "")[:35],
            ])

        pdf_bytes = _build_pdf_report(
            "Hostel Nightly Attendance & Roll Call",
            headers,
            data,
            orientation="portrait",
        )
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="Hostel_Attendance_Report.pdf"'
        return response


# ==============================================================================
# 5. FEES & DUES REPORT
# ==============================================================================
class ExportHostelFeesExcelView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        records = FeeRecord.objects.filter(fee_type="HOSTEL").select_related("student").prefetch_related("payments")
        hostel_id = request.GET.get("hostel")
        if hostel_id and hostel_id.isdigit():
            student_ids = HostelAllocation.objects.filter(hostel_id=int(hostel_id)).values_list("student_id", flat=True)
            records = records.filter(student_id__in=student_ids)

        records = records.order_by("-due_date", "student__name")

        wb = Workbook()
        ws = wb.active
        ws.title = "Hostel Fees"

        headers = [
            "Student Name",
            "Roll No",
            "Description / Subtype",
            "Total Amount (INR)",
            "Paid Amount (INR)",
            "Balance Due (INR)",
            "Due Date",
            "Status",
        ]
        _apply_excel_header_style(ws, headers)

        for f in records:
            paid = sum((p.amount for p in f.payments.all()), Decimal("0.00"))
            bal = max(Decimal("0.00"), f.amount - paid)
            ws.append([
                f.student.name,
                f.student.roll_no,
                f.description or "Hostel Fee",
                float(f.amount),
                float(paid),
                float(bal),
                f.due_date.strftime("%d-%m-%Y") if f.due_date else "—",
                f.get_status_display(),
            ])

        _autofit_excel_columns(ws)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="Hostel_Fee_Collections.xlsx"'
        wb.save(response)
        return response


class ExportHostelFeesPdfView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        records = FeeRecord.objects.filter(fee_type="HOSTEL").select_related("student").prefetch_related("payments")
        hostel_id = request.GET.get("hostel")
        if hostel_id and hostel_id.isdigit():
            student_ids = HostelAllocation.objects.filter(hostel_id=int(hostel_id)).values_list("student_id", flat=True)
            records = records.filter(student_id__in=student_ids)

        records = records.order_by("-due_date", "student__name")

        headers = ["Student Name", "Roll No", "Fee Item", "Total", "Paid", "Balance Due", "Due Date", "Status"]
        data = []
        for f in records:
            paid = sum((p.amount for p in f.payments.all()), Decimal("0.00"))
            bal = max(Decimal("0.00"), f.amount - paid)
            data.append([
                f.student.name,
                f.student.roll_no,
                (f.description or "Hostel Fee")[:24],
                f"Rs {f.amount:.0f}",
                f"Rs {paid:.0f}",
                f"Rs {bal:.0f}",
                f.due_date.strftime("%d-%b-%y") if f.due_date else "—",
                f.get_status_display(),
            ])

        pdf_bytes = _build_pdf_report(
            "Hostel Fees & Outstanding Dues Summary",
            headers,
            data,
            orientation="portrait",
        )
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="Hostel_Fee_Collections.pdf"'
        return response


# ==============================================================================
# 6. COMPLAINTS & MAINTENANCE REPORT
# ==============================================================================
class ExportHostelComplaintsExcelView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        complaints = HostelComplaint.objects.filter(hostel_id__in=assigned_ids).select_related(
            "student", "hostel", "room", "assigned_staff"
        )
        if hostel_id and hostel_id.isdigit():
            complaints = complaints.filter(hostel_id=int(hostel_id))

        complaints = complaints.order_by("-created_at")

        wb = Workbook()
        ws = wb.active
        ws.title = "Hostel Complaints"

        headers = [
            "Ticket ID",
            "Hostel Code",
            "Room No",
            "Student Name",
            "Roll No",
            "Category",
            "Title",
            "Priority",
            "Status",
            "Logged Date",
            "Assigned Staff",
            "Resolution",
        ]
        _apply_excel_header_style(ws, headers)

        for c in complaints:
            staff = c.assigned_staff.get_full_name() or c.assigned_staff.username if c.assigned_staff else "Unassigned"
            ws.append([
                c.complaint_id,
                c.hostel.code,
                c.room.room_number if c.room else "—",
                c.student.name,
                c.student.roll_no,
                c.get_category_display(),
                c.title,
                c.get_priority_display(),
                c.get_status_display(),
                c.created_at.strftime("%d-%m-%Y"),
                staff,
                c.resolution or "—",
            ])

        _autofit_excel_columns(ws)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="Hostel_Maintenance_Complaints.xlsx"'
        wb.save(response)
        return response


class ExportHostelComplaintsPdfView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        complaints = HostelComplaint.objects.filter(hostel_id__in=assigned_ids).select_related(
            "student", "hostel", "room", "assigned_staff"
        )
        if hostel_id and hostel_id.isdigit():
            complaints = complaints.filter(hostel_id=int(hostel_id))

        complaints = complaints.order_by("-created_at")

        headers = ["Ticket", "Hostel", "Room", "Student", "Category", "Priority", "Status", "Date", "Staff Assigned"]
        data = []
        for c in complaints:
            staff = c.assigned_staff.get_full_name() or c.assigned_staff.username if c.assigned_staff else "—"
            data.append([
                c.complaint_id,
                c.hostel.code,
                c.room.room_number if c.room else "—",
                c.student.name[:16],
                c.get_category_display(),
                c.get_priority_display(),
                c.get_status_display(),
                c.created_at.strftime("%d-%b-%y"),
                staff[:16],
            ])

        pdf_bytes = _build_pdf_report(
            "Hostel Maintenance & Repair Tickets",
            headers,
            data,
            orientation="landscape",
        )
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="Hostel_Maintenance_Complaints.pdf"'
        return response


# ==============================================================================
# 7. VISITOR PASSES REPORT
# ==============================================================================
class ExportHostelVisitorsExcelView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        visitors = HostelVisitor.objects.filter(hostel_id__in=assigned_ids).select_related(
            "student", "hostel", "approved_by"
        )
        if hostel_id and hostel_id.isdigit():
            visitors = visitors.filter(hostel_id=int(hostel_id))

        visitors = visitors.order_by("-visit_date", "-created_at")

        wb = Workbook()
        ws = wb.active
        ws.title = "Visitor Gate Logs"

        headers = [
            "Visit Date",
            "Visitor Name",
            "Relationship",
            "Contact Phone",
            "Resident Student",
            "Student Roll No",
            "Hostel Code",
            "Purpose",
            "Entry Time",
            "Exit Time",
            "Status",
            "Approved By",
            "Remarks",
        ]
        _apply_excel_header_style(ws, headers)

        for v in visitors:
            approver = v.approved_by.get_full_name() or v.approved_by.username if v.approved_by else "—"
            ws.append([
                v.visit_date.strftime("%d-%m-%Y"),
                v.visitor_name,
                v.relationship,
                v.phone,
                v.student.name,
                v.student.roll_no,
                v.hostel.code,
                v.purpose,
                v.entry_time.strftime("%H:%M") if v.entry_time else "—",
                v.exit_time.strftime("%H:%M") if v.exit_time else "—",
                v.get_status_display(),
                approver,
                v.remarks or "—",
            ])

        _autofit_excel_columns(ws)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="Hostel_Visitor_Gate_Passes.xlsx"'
        wb.save(response)
        return response


class ExportHostelVisitorsPdfView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        visitors = HostelVisitor.objects.filter(hostel_id__in=assigned_ids).select_related(
            "student", "hostel", "approved_by"
        )
        if hostel_id and hostel_id.isdigit():
            visitors = visitors.filter(hostel_id=int(hostel_id))

        visitors = visitors.order_by("-visit_date", "-created_at")

        headers = ["Date", "Visitor Name", "Relation", "Phone", "Student", "Hostel", "Purpose", "In/Out", "Status", "Approval"]
        data = []
        for v in visitors:
            approver = v.approved_by.get_full_name() or v.approved_by.username if v.approved_by else "—"
            in_out = f"{v.entry_time.strftime('%H:%M') if v.entry_time else '—'} / {v.exit_time.strftime('%H:%M') if v.exit_time else '—'}"
            data.append([
                v.visit_date.strftime("%d-%b-%y"),
                v.visitor_name,
                v.relationship,
                v.phone,
                v.student.name[:16],
                v.hostel.code,
                v.purpose[:20],
                in_out,
                v.get_status_display(),
                approver[:14],
            ])

        pdf_bytes = _build_pdf_report(
            "Hostel Visitor Security & Gate Passes",
            headers,
            data,
            orientation="landscape",
        )
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="Hostel_Visitor_Gate_Passes.pdf"'
        return response


# ==============================================================================
# 8. APPLICATION REPORT
# ==============================================================================
class ExportHostelApplicationsExcelView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        apps = HostelApplication.objects.filter(hostel_preference_id__in=assigned_ids).select_related(
            "student", "hostel_preference", "reviewed_by"
        )
        if hostel_id and hostel_id.isdigit():
            apps = apps.filter(hostel_preference_id=int(hostel_id))

        apps = apps.order_by("-application_date", "-created_at")

        wb = Workbook()
        ws = wb.active
        ws.title = "Applications Log"

        headers = [
            "Application ID",
            "Date",
            "Student Name",
            "Roll No",
            "Preferred Hostel",
            "Room Preference",
            "Academic Year",
            "Status",
            "Reviewed By",
            "Reviewed Date",
            "Reason / Notes",
        ]
        _apply_excel_header_style(ws, headers)

        for a in apps:
            reviewer = a.reviewed_by.get_full_name() or a.reviewed_by.username if a.reviewed_by else "—"
            rev_date = a.reviewed_date.strftime("%d-%m-%Y") if a.reviewed_date else "—"
            ws.append([
                a.application_id,
                a.application_date.strftime("%d-%m-%Y"),
                a.student.name,
                a.student.roll_no,
                a.hostel_preference.code,
                a.get_room_type_preference_display(),
                a.academic_year,
                a.get_status_display(),
                reviewer,
                rev_date,
                a.rejection_reason or a.reason or "—",
            ])

        _autofit_excel_columns(ws)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="Hostel_Applications_Report.xlsx"'
        wb.save(response)
        return response


class ExportHostelApplicationsPdfView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        apps = HostelApplication.objects.filter(hostel_preference_id__in=assigned_ids).select_related(
            "student", "hostel_preference", "reviewed_by"
        )
        if hostel_id and hostel_id.isdigit():
            apps = apps.filter(hostel_preference_id=int(hostel_id))

        apps = apps.order_by("-application_date", "-created_at")

        headers = ["App ID", "Date", "Student", "Roll No", "Hostel", "Room Type", "Status", "Reviewer"]
        data = []
        for a in apps:
            reviewer = a.reviewed_by.get_full_name() or a.reviewed_by.username if a.reviewed_by else "—"
            data.append([
                a.application_id,
                a.application_date.strftime("%d-%b-%y"),
                a.student.name[:18],
                a.student.roll_no,
                a.hostel_preference.code,
                a.get_room_type_preference_display()[:6],
                a.get_status_display(),
                reviewer[:14],
            ])

        pdf_bytes = _build_pdf_report(
            "Hostel Admission Applications Report",
            headers,
            data,
            orientation="landscape",
        )
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="Hostel_Applications_Report.pdf"'
        return response


# ==============================================================================
# 9. CHECK-IN / CHECK-OUT LIFECYCLE REPORT
# ==============================================================================
class ExportHostelCheckInOutExcelView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        events = HostelCheckInOut.objects.filter(allocation__hostel_id__in=assigned_ids).select_related(
            "student", "allocation__hostel", "allocation__room", "allocation__bed", "processed_by"
        )
        if hostel_id and hostel_id.isdigit():
            events = events.filter(allocation__hostel_id=int(hostel_id))

        events = events.order_by("-event_date", "-event_time")

        wb = Workbook()
        ws = wb.active
        ws.title = "Lifecycle Events"

        headers = [
            "Event Type",
            "Date",
            "Time",
            "Student Name",
            "Roll No",
            "Hostel",
            "Room No",
            "Bed No",
            "Processed By",
            "Reason / Remarks",
        ]
        _apply_excel_header_style(ws, headers)

        for ev in events:
            proc = ev.processed_by.get_full_name() or ev.processed_by.username if ev.processed_by else "—"
            ws.append([
                ev.get_event_type_display(),
                ev.event_date.strftime("%d-%m-%Y"),
                ev.event_time.strftime("%H:%M") if ev.event_time else "—",
                ev.student.name,
                ev.student.roll_no,
                ev.allocation.hostel.code,
                ev.allocation.room.room_number,
                ev.allocation.bed.bed_number,
                proc,
                ev.remarks or ev.reason or "—",
            ])

        _autofit_excel_columns(ws)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="Hostel_CheckInOut_Report.xlsx"'
        wb.save(response)
        return response


class ExportHostelCheckInOutPdfView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        events = HostelCheckInOut.objects.filter(allocation__hostel_id__in=assigned_ids).select_related(
            "student", "allocation__hostel", "allocation__room", "allocation__bed", "processed_by"
        )
        if hostel_id and hostel_id.isdigit():
            events = events.filter(allocation__hostel_id=int(hostel_id))

        events = events.order_by("-event_date", "-event_time")

        headers = ["Event", "Date", "Time", "Student", "Roll No", "Hostel", "Room", "Bed", "Processed By"]
        data = []
        for ev in events:
            proc = ev.processed_by.get_full_name() or ev.processed_by.username if ev.processed_by else "—"
            data.append([
                ev.get_event_type_display(),
                ev.event_date.strftime("%d-%b-%y"),
                ev.event_time.strftime("%H:%M") if ev.event_time else "—",
                ev.student.name[:18],
                ev.student.roll_no,
                ev.allocation.hostel.code,
                ev.allocation.room.room_number,
                ev.allocation.bed.bed_number,
                proc[:14],
            ])

        pdf_bytes = _build_pdf_report(
            "Hostel Check-in & Check-out Lifecycle Report",
            headers,
            data,
            orientation="landscape",
        )
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="Hostel_CheckInOut_Report.pdf"'
        return response


# ==============================================================================
# 10. ROOM TRANSFER HISTORY REPORT
# ==============================================================================
class ExportHostelTransfersExcelView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        transfers = HostelRoomTransfer.objects.filter(
            Q(old_hostel_id__in=assigned_ids) | Q(new_hostel_id__in=assigned_ids)
        ).select_related(
            "student", "old_hostel", "old_room", "old_bed", "new_hostel", "new_room", "new_bed", "approved_by"
        )
        if hostel_id and hostel_id.isdigit():
            hid = int(hostel_id)
            transfers = transfers.filter(Q(old_hostel_id=hid) | Q(new_hostel_id=hid))

        transfers = transfers.order_by("-transfer_date", "-created_at")

        wb = Workbook()
        ws = wb.active
        ws.title = "Transfer History"

        headers = [
            "Transfer Date",
            "Student Name",
            "Roll No",
            "From Hostel",
            "From Room/Bed",
            "To Hostel",
            "To Room/Bed",
            "Reason",
            "Approved By",
            "Remarks",
        ]
        _apply_excel_header_style(ws, headers)

        for t in transfers:
            approver = t.approved_by.get_full_name() or t.approved_by.username if t.approved_by else "—"
            ws.append([
                t.transfer_date.strftime("%d-%m-%Y"),
                t.student.name,
                t.student.roll_no,
                t.old_hostel.code,
                f"R{t.old_room.room_number} / B{t.old_bed.bed_number}",
                t.new_hostel.code,
                f"R{t.new_room.room_number} / B{t.new_bed.bed_number}",
                t.reason,
                approver,
                t.remarks or "—",
            ])

        _autofit_excel_columns(ws)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="Hostel_Transfer_History.xlsx"'
        wb.save(response)
        return response


class ExportHostelTransfersPdfView(LoginRequiredMixin, View):
    def get(self, request):
        if not _is_hostel_admin(request.user):
            raise PermissionDenied("Hostel management access required.")

        assigned_ids = get_user_assigned_hostel_ids(request.user)
        hostel_id = request.GET.get("hostel")
        transfers = HostelRoomTransfer.objects.filter(
            Q(old_hostel_id__in=assigned_ids) | Q(new_hostel_id__in=assigned_ids)
        ).select_related(
            "student", "old_hostel", "old_room", "old_bed", "new_hostel", "new_room", "new_bed", "approved_by"
        )
        if hostel_id and hostel_id.isdigit():
            hid = int(hostel_id)
            transfers = transfers.filter(Q(old_hostel_id=hid) | Q(new_hostel_id=hid))

        transfers = transfers.order_by("-transfer_date", "-created_at")

        headers = ["Date", "Student", "Roll No", "From (Hostel/Room/Bed)", "To (Hostel/Room/Bed)", "Reason", "Approved By"]
        data = []
        for t in transfers:
            approver = t.approved_by.get_full_name() or t.approved_by.username if t.approved_by else "—"
            from_loc = f"{t.old_hostel.code} R{t.old_room.room_number}-B{t.old_bed.bed_number}"
            to_loc = f"{t.new_hostel.code} R{t.new_room.room_number}-B{t.new_bed.bed_number}"
            data.append([
                t.transfer_date.strftime("%d-%b-%y"),
                t.student.name[:16],
                t.student.roll_no,
                from_loc,
                to_loc,
                t.reason[:24],
                approver[:14],
            ])

        pdf_bytes = _build_pdf_report(
            "Hostel Room Transfer Audit History",
            headers,
            data,
            orientation="landscape",
        )
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="Hostel_Transfer_History.pdf"'
        return response

