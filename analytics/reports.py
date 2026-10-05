import io
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from django.http import HttpResponse
from django.utils import timezone

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

from . import services


def generate_analytics_excel(scope, export_type="overall"):
    """
    Generates a professionally formatted Excel spreadsheet using openpyxl.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Analytics Overview"

    # Styling definitions
    brand_blue = "1D4ED8"
    accent_blue = "DBEAFE"
    header_fill = PatternFill(start_color=brand_blue, end_color=brand_blue, fill_type="solid")
    sub_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    kpi_fill = PatternFill(start_color=accent_blue, end_color=accent_blue, fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    font_title = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    font_sub = Font(name="Arial", size=10, bold=False, color="FFFFFF")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Arial", size=10, bold=True, color="000000")
    font_norm = Font(name="Arial", size=9, bold=False, color="000000")
    font_muted = Font(name="Arial", size=8, italic=True, color="64748B")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # 1. Title Block
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "SMART COLLEGE — TAKSHASHILA UNIVERSITY"
    title_cell.font = font_title
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:H2")
    sub_cell = ws["A2"]
    sub_cell.value = f"AI & Advanced Analytics Report ({export_type.upper()})"
    sub_cell.font = font_sub
    sub_cell.fill = sub_fill
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # 2. Metadata Block
    school_name = scope["school"].name if scope["school"] else "All Schools / Institutional"
    course_name = scope["course"].name if scope["course"] else "All Courses"
    year_str = f"Year {scope['year']}" if scope["year"] else "All Years"

    ws.cell(row=4, column=1, value="Generated On:").font = font_bold
    ws.cell(row=4, column=2, value=timezone.now().strftime("%d %b %Y, %I:%M %p")).font = font_norm
    ws.cell(row=4, column=4, value="School / Dept:").font = font_bold
    ws.cell(row=4, column=5, value=school_name).font = font_norm

    ws.cell(row=5, column=1, value="Degree / Course:").font = font_bold
    ws.cell(row=5, column=2, value=course_name).font = font_norm
    ws.cell(row=5, column=4, value="Class / Year:").font = font_bold
    ws.cell(row=5, column=5, value=year_str).font = font_norm

    kpi = services.get_kpi_summary(scope)

    # 3. KPI Summary Cards in Excel
    ws.cell(row=7, column=1, value="INSTITUTIONAL PERFORMANCE SUMMARY").font = font_bold
    headers_kpi = ["Total Students", "Total Faculty", "Average Attendance", "Average Marks", "Total Fees", "Collected", "Outstanding", "Placement Rate"]
    values_kpi = [
        kpi["total_students"],
        kpi["total_faculty"],
        f"{kpi['avg_attendance']}%" if kpi["avg_attendance"] is not None else "—",
        f"{kpi['avg_marks']}%" if kpi["avg_marks"] is not None else "—",
        f"₹{kpi['fees_total']:,.2f}",
        f"₹{kpi['fees_collected']:,.2f}",
        f"₹{kpi['fees_outstanding']:,.2f}",
        f"{kpi['placement_rate']}%",
    ]

    for col_idx, (h, v) in enumerate(zip(headers_kpi, values_kpi), start=1):
        c_h = ws.cell(row=8, column=col_idx, value=h)
        c_h.font = font_bold
        c_h.fill = kpi_fill
        c_h.border = thin_border
        c_h.alignment = Alignment(horizontal="center", vertical="center")

        c_v = ws.cell(row=9, column=col_idx, value=v)
        c_v.font = font_bold
        c_v.border = thin_border
        c_v.alignment = Alignment(horizontal="center", vertical="center")

    curr_row = 11

    # 4. Data Table according to export_type
    if export_type in ["risk", "overall"]:
        risk_data = services.get_risk_student_analytics(scope)
        ws.cell(row=curr_row, column=1, value="ACADEMIC RISK INDICATOR & SUPPORT RECOMMENDATIONS").font = font_bold
        curr_row += 1

        headers = ["Roll No", "Student Name", "Course", "Year", "Attendance %", "Marks %", "Fee Balance", "Risk Score", "Risk Level", "Flagged Reasons"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        curr_row += 1

        for r_idx, item in enumerate(risk_data["students"]):
            row_fill = alt_fill if r_idx % 2 == 1 else PatternFill(fill_type=None)
            reasons_text = "; ".join(item["reasons"])
            values = [
                item["roll_no"],
                item["name"],
                item["course"],
                f"Yr {item['year']}",
                f"{item['attendance_pct']}%" if item["attendance_pct"] is not None else "No Data",
                f"{item['marks_pct']}%" if item["marks_pct"] is not None else "No Data",
                f"₹{item['fee_balance']:,.2f}",
                f"{item['risk_score']}/100",
                item["risk_level"],
                reasons_text,
            ]
            for col_idx, val in enumerate(values, start=1):
                c = ws.cell(row=curr_row, column=col_idx, value=val)
                c.font = font_norm
                c.border = thin_border
                c.fill = row_fill
                if col_idx in [1, 4, 5, 6, 7, 8, 9]:
                    c.alignment = Alignment(horizontal="center")
            curr_row += 1

        curr_row += 2

    if export_type in ["attendance", "overall"]:
        att_data = services.get_attendance_analytics(scope)
        ws.cell(row=curr_row, column=1, value="ATTENDANCE ANALYSIS & SHORTAGE LIST").font = font_bold
        curr_row += 1

        headers = ["Roll No", "Student Name", "Course", "Year", "Total Classes", "Present", "Absent", "Attendance %", "Category"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = sub_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
        curr_row += 1

        for r_idx, item in enumerate(att_data["risk_list"]):
            row_fill = alt_fill if r_idx % 2 == 1 else PatternFill(fill_type=None)
            values = [
                item["roll_no"],
                item["name"],
                item["course"],
                f"Yr {item['year']}",
                item["total_classes"],
                item["present_count"],
                item["absent_count"],
                f"{item['attendance_pct']}%" if item["attendance_pct"] is not None else "No Data",
                item["risk_level"],
            ]
            for col_idx, val in enumerate(values, start=1):
                c = ws.cell(row=curr_row, column=col_idx, value=val)
                c.font = font_norm
                c.border = thin_border
                c.fill = row_fill
                if col_idx in [1, 4, 5, 6, 7, 8, 9]:
                    c.alignment = Alignment(horizontal="center")
            curr_row += 1

        curr_row += 2

    if export_type in ["academic", "overall"]:
        acad_data = services.get_academic_analytics(scope)
        ws.cell(row=curr_row, column=1, value="SUBJECT PERFORMANCE METRICS").font = font_bold
        curr_row += 1

        headers = ["Subject Name", "Code", "Faculty", "Evaluated Students", "Average Marks %", "Pass Rate %", "Highest %", "Lowest %"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
        curr_row += 1

        for r_idx, item in enumerate(acad_data["subjects"]):
            row_fill = alt_fill if r_idx % 2 == 1 else PatternFill(fill_type=None)
            values = [
                item["name"],
                item["code"],
                item["faculty_name"],
                item["students_count"],
                f"{item['avg_percentage']}%",
                f"{item['pass_rate']}%",
                f"{item['max_score']}%",
                f"{item['min_score']}%",
            ]
            for col_idx, val in enumerate(values, start=1):
                c = ws.cell(row=curr_row, column=col_idx, value=val)
                c.font = font_norm
                c.border = thin_border
                c.fill = row_fill
                if col_idx in [2, 4, 5, 6, 7, 8]:
                    c.alignment = Alignment(horizontal="center")
            curr_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and len(val_str) < 60:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"SmartCollege_Analytics_{export_type}_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def generate_analytics_pdf(scope, export_type="overall"):
    """
    Generates a clean PDF document using ReportLab.
    """
    response = HttpResponse(content_type="application/pdf")
    filename = f"SmartCollege_Analytics_Report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    response["Content-Disposition"] = f'inline; filename="{filename}"'

    doc = SimpleDocTemplate(
        response,
        pagesize=A4,
        rightMargin=26,
        leftMargin=26,
        topMargin=26,
        bottomMargin=26
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#1D4ED8"),
        alignment=TA_CENTER
    )
    sub_style = ParagraphStyle(
        "DocSub",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#334155"),
        alignment=TA_CENTER
    )
    meta_style = ParagraphStyle(
        "MetaText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#64748B")
    )
    h2_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#1D4ED8"),
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10
    )
    insight_style = ParagraphStyle(
        "InsightText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#1E293B")
    )

    elements = []

    # Title Block
    elements.append(Paragraph("SMART COLLEGE · TAKSHASHILA UNIVERSITY", title_style))
    elements.append(Spacer(1, 3))
    elements.append(Paragraph(f"AI & ADVANCED ANALYTICS REPORT · {export_type.upper()}", sub_style))
    elements.append(Spacer(1, 8))

    # Metadata & Filters
    school_name = scope["school"].name if scope["school"] else "Institutional (All Schools)"
    course_name = scope["course"].name if scope["course"] else "All Courses"
    year_str = f"Year {scope['year']}" if scope["year"] else "All Years"
    gen_time = timezone.now().strftime("%d %B %Y, %I:%M %p")

    meta_table_data = [
        [
            Paragraph(f"<b>Generated:</b> {gen_time}", meta_style),
            Paragraph(f"<b>Scope/Dept:</b> {school_name}", meta_style),
        ],
        [
            Paragraph(f"<b>Course:</b> {course_name}", meta_style),
            Paragraph(f"<b>Class:</b> {year_str}", meta_style),
        ]
    ]
    meta_table = Table(meta_table_data, colWidths=[270, 270])
    meta_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ("PADDING", (0, 0), (-1, -1), 4),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 10))

    # KPI Summary
    kpi = services.get_kpi_summary(scope)
    elements.append(Paragraph("Executive KPI Summary", h2_style))

    kpi_data = [
        ["Total Students", "Total Faculty", "Average Attendance", "Average Marks"],
        [
            str(kpi["total_students"]),
            str(kpi["total_faculty"]),
            f"{kpi['avg_attendance']}%" if kpi["avg_attendance"] is not None else "—",
            f"{kpi['avg_marks']}%" if kpi["avg_marks"] is not None else "—",
        ],
        ["Total Fees", "Fees Collected", "Outstanding Balance", "Placement Rate"],
        [
            f"INR {kpi['fees_total']:,.0f}",
            f"INR {kpi['fees_collected']:,.0f}",
            f"INR {kpi['fees_outstanding']:,.0f}",
            f"{kpi['placement_rate']}%",
        ]
    ]
    t_kpi = Table(kpi_data, colWidths=[135, 135, 135, 135])
    t_kpi.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8.5),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#EFF6FF")),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 1), (-1, 1), 10),
        ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#334155")),
        ("TEXTCOLOR", (0, 2), (-1, 2), colors.white),
        ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"),
        ("FONTSIZE", (0, 2), (-1, 2), 8.5),
        ("BACKGROUND", (0, 3), (-1, 3), colors.HexColor("#F8FAFC")),
        ("FONTNAME", (0, 3), (-1, 3), "Helvetica-Bold"),
        ("FONTSIZE", (0, 3), (-1, 3), 10),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(t_kpi)
    elements.append(Spacer(1, 10))

    # AI Academic Insights
    academic_data = services.get_academic_analytics(scope)
    att_data = services.get_attendance_analytics(scope)
    risk_data = services.get_risk_student_analytics(scope)
    insights = services.get_ai_insights(scope, kpi, academic_data, att_data, risk_data)

    if insights:
        elements.append(Paragraph("AI Academic Insights & Diagnostic Observations", h2_style))
        insight_rows = []
        for ins in insights[:4]:
            text = f"<b>{ins['title']}:</b> {ins['message']} <i>[{ins['metric']}]</i>"
            insight_rows.append([Paragraph(text, insight_style)])
        t_ins = Table(insight_rows, colWidths=[540])
        t_ins.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("PADDING", (0, 0), (-1, -1), 5),
            ("LINEBELOW", (0, 0), (-1, -2), 0.5, colors.HexColor("#E2E8F0")),
        ]))
        elements.append(t_ins)
        elements.append(Spacer(1, 10))

    # Academic Risk Students Table
    if risk_data["students"]:
        elements.append(Paragraph("Academic Risk Indicator (Support Recommendation)", h2_style))
        risk_table_rows = [
            ["Roll No", "Student Name", "Course", "Yr", "Att %", "Marks %", "Risk", "Flagged Reasons"]
        ]
        for item in risk_data["students"][:15]:
            reasons_str = "; ".join(item["reasons"])
            risk_table_rows.append([
                Paragraph(item["roll_no"], body_style),
                Paragraph(item["name"], body_style),
                Paragraph(item["course"], body_style),
                Paragraph(str(item["year"]), body_style),
                Paragraph(f"{item['attendance_pct']}%" if item["attendance_pct"] is not None else "—", body_style),
                Paragraph(f"{item['marks_pct']}%" if item["marks_pct"] is not None else "—", body_style),
                Paragraph(f"<b>{item['risk_level']}</b>", body_style),
                Paragraph(reasons_str, body_style),
            ])

        t_risk = Table(risk_table_rows, colWidths=[55, 85, 95, 25, 40, 45, 45, 150])
        t_risk.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 7.5),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("PADDING", (0, 0), (-1, -1), 3),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ]))
        elements.append(t_risk)

    elements.append(Spacer(1, 12))
    elements.append(Paragraph(
        "<i>Disclaimer: This document is an Academic Support Recommendation generated strictly from internal institutional ERP data to guide academic mentoring and administrative support.</i>",
        meta_style
    ))

    doc.build(elements)
    return response
